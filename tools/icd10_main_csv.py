# -*- coding: utf-8 -*-
"""国标《疾病分类与代码》(GB/T 14396-2016) → 主要诊断主表 CSV。

从 Excel Sheet1 读取：
  A列=主要编码（剥除尾部 †：† 仅标注“该病因码可作主要诊断”，码本身仍是主诊断可用码）
  C列=疾病名称（清洗 †/*、换行、多余空白）
输出 data/terms-inbox/icd10_main.csv（code,name，UTF-8-sig），保留 Excel 原始顺序。

B列=附加编码（* 表现码，临床不能单独作主要诊断）本次不输出。

用法: python3 tools/icd10_main_csv.py
"""
import csv
import os
import re

from openpyxl import load_workbook

SRC = os.path.expanduser("~/Downloads/《疾病分类与代码》（GBT 14396-2016）国标版.xlsx")
OUT = os.path.join(os.path.dirname(__file__), "..", "data", "terms-inbox", "icd10_main.csv")

WS = "Sheet1"


def clean_code(raw: str) -> str:
    """主要编码：去尾随 †/* 与空白。"""
    return (raw or "").strip().rstrip("†*").strip()


def clean_name(raw) -> str:
    """疾病名称：去 †/*、换行制表、压缩多余空白。"""
    if raw is None:
        return ""
    s = str(raw)
    s = s.replace("†", "").replace("*", "")
    s = re.sub(r"[\r\n\t]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def main():
    wb = load_workbook(SRC, read_only=True, data_only=True)
    if WS not in wb.sheetnames:
        raise SystemExit(f"未找到工作表 {WS}，实际: {wb.sheetnames}")
    ws = wb[WS]

    # 跳过表头（第 1 行），A 列非空才输出
    rows = []
    stripped = 0  # 剥除 † 的行数
    for i, row in enumerate(ws.iter_rows(min_row=2, max_col=3, values_only=True), start=2):
        a, _b, c = (row + (None,) * 3)[:3]
        raw_a = (a or "").strip()
        if not raw_a:
            continue
        code = clean_code(raw_a)
        name = clean_name(c)
        if raw_a.rstrip("†*").strip() != raw_a:  # 原始含尾符号（†/*）
            stripped += 1
        rows.append((code, name, i))

    # 重复 code 检查（同码多名称时保留首条并提示）
    seen, dup = set(), []
    uniq = []
    for code, name, line in rows:
        if code in seen:
            dup.append((code, line))
            continue
        seen.add(code)
        uniq.append((code, name))

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["code", "name"])
        w.writerows(uniq)

    print(f"源 Excel 行(A非空): {len(rows)}")
    print(f"剥除 †/* 行: {stripped}")
    print(f"输出主表: {len(uniq)} 行 -> {OUT}")
    if dup:
        print(f"重复 code {len(dup)} 条(保留首条)，示例: {dup[:5]}")


if __name__ == "__main__":
    main()
